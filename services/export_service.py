from docx import Document
from typing import Dict, Any

def generate_doc(itinerary_data: Dict[str, Any]) -> str:
    """
    Generates a Word document from itinerary data.
    
    Args:
        itinerary_data: Expected to have 'destination', 'itinerary' (list), and 'summary'.
        
    Returns:
        The path to the generated .docx file.
    """
    doc = Document()
    
    doc.add_heading('Travel Itinerary', 0)
    
    destination = itinerary_data.get("destination", "Your Trip")
    doc.add_paragraph(f"Destination: {destination}")
    doc.add_paragraph(f"Summary: {itinerary_data.get('summary', 'Enjoy your travels!')}")
    
    doc.add_heading('Day 1', level=1)
    
    itinerary_items = itinerary_data.get("itinerary", [])
    for item in itinerary_items:
        time = item.get("time", "")
        place = item.get("place", "")
        doc.add_paragraph(f"- {time}: {place}", style='ListBullet')
        
    file_path = "itinerary.docx"
    doc.save(file_path)
    return file_path
